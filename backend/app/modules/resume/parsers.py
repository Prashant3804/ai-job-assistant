import io
import os
import re
from abc import ABC, abstractmethod
from typing import Optional
from fastapi import HTTPException, status

class BaseResumeParser(ABC):
    @abstractmethod
    def extract_text(self, file_bytes: bytes) -> str:
        pass

class PDFResumeParser(BaseResumeParser):
    def extract_text(self, file_bytes: bytes) -> str:
        if not file_bytes:
            raise ValueError("Empty PDF file bytes provided.")

        # Try PyMuPDF (fitz) first
        try:
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            full_text = "\n".join(text_parts).strip()
            if full_text:
                return full_text
        except Exception:
            pass

        # Fallback to pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
            full_text = "\n".join(text_parts).strip()
            if full_text:
                return full_text
        except Exception as e:
            raise ValueError(f"Failed to extract text from PDF document: {str(e)}")

        raise ValueError("No readable text found in PDF document (may be scanned image).")

class DocxResumeParser(BaseResumeParser):
    def extract_text(self, file_bytes: bytes) -> str:
        if not file_bytes:
            raise ValueError("Empty DOCX file bytes provided.")

        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            full_text = "\n".join(text_parts).strip()
            if not full_text:
                raise ValueError("DOCX document contains no text paragraphs or tables.")
            return full_text
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

class PlainTextResumeParser(BaseResumeParser):
    def extract_text(self, file_bytes: bytes) -> str:
        if not file_bytes:
            raise ValueError("Empty text content provided.")

        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                return file_bytes.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode text file with standard encodings.")

def get_parser_for_file(filename: str) -> BaseResumeParser:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return PDFResumeParser()
    elif ext in [".docx", ".doc"]:
        return DocxResumeParser()
    elif ext in [".txt", ".md", ".text"]:
        return PlainTextResumeParser()
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Only PDF (.pdf), Word (.docx), and Plain Text (.txt) are supported."
        )

def validate_resume_file(filename: str, file_bytes: bytes, max_size_mb: int = 10) -> None:
    if not file_bytes or len(file_bytes) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty."
        )

    max_bytes = max_size_mb * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise HTTPException(
            status_code=getattr(status, 'HTTP_413_CONTENT_TOO_LARGE', 413),
            detail=f"File size ({len(file_bytes) / (1024*1024):.1f}MB) exceeds maximum allowed limit of {max_size_mb}MB."
        )

    ext = os.path.splitext(filename)[1].lower()
    if ext not in [".pdf", ".docx", ".doc", ".txt", ".md"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file extension '{ext}'. Only PDF, DOCX, and TXT files are accepted."
        )

    # Basic header signature checks
    if ext == ".pdf" and not file_bytes.startswith(b"%PDF"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid PDF file: Missing %PDF signature."
        )
    elif ext == ".docx" and not file_bytes.startswith(b"PK\x03\x04"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid DOCX file: Missing standard OpenXML package header."
        )
